"""Régressions B7/B8/B9 : vrais formulaires et données exclusivement fictives."""
import re


def ajouter_absence(client, uid, debut='2026-09-07', fin=None, motif='Congé payé'):
    return client.post('/absences', data={
        'user_id': uid, 'motif': motif, 'date_debut': debut,
        'date_fin': fin or debut,
    }, follow_redirects=True)


def traiter(client, uid, reference=None):
    if reference is None:
        html = client.get('/prepa_paie?mois=9&annee=2026').get_data(as_text=True)
        match = re.search(r'name="reference_' + str(uid) + r'" value="([^"]+)"', html)
        reference = match.group(1) if match else ''
    return client.post('/prepa_paie/traiter', data={
        'mois': 9, 'annee': 2026, 'user_ids': [uid], f'traite_{uid}': '1',
        f'reference_{uid}': reference,
    }, follow_redirects=True)


def demande_partielle(db, uid):
    row = db.execute("""INSERT INTO demandes_recup
        (user_id, date_debut, date_fin, nb_jours, nb_heures, type_demande,
         heure_debut, heure_fin, statut)
        VALUES (?, '2026-09-07', '2026-09-07', .43, 3, 'partielle',
                '14:00', '17:00', 'en_attente_direction')""", (uid,))
    db.commit()
    return row.lastrowid


def approuver(client, did, type_demande='recup'):
    return client.post('/validation_demandes_recup', data={
        'demande_id': did, 'action': 'valider', 'demande_type': type_demande,
    }, follow_redirects=True)


def test_b7_double_cp_refuse(admin_client, db, sample_users):
    uid = sample_users['salarie_id']
    ajouter_absence(admin_client, uid)
    ajouter_absence(admin_client, uid)
    assert db.execute('SELECT COUNT(*) FROM absences WHERE user_id=?', (uid,)).fetchone()[0] == 1
    assert db.execute('SELECT cp_pris FROM users WHERE id=?', (uid,)).fetchone()[0] == 1


def test_b8_absence_apres_verification(admin_client, db, sample_users, sample_contrat):
    uid = sample_users['salarie_id']
    traiter(admin_client, uid)
    assert db.execute('SELECT traite FROM prepa_paie_statut WHERE user_id=?', (uid,)).fetchone()[0] == 1
    ajouter_absence(admin_client, uid, motif='Arrêt maladie')
    html = admin_client.get('/prepa_paie?mois=9&annee=2026').get_data(as_text=True)
    assert 'Modifié depuis la dernière vérification' in html
    assert db.execute('SELECT traite FROM prepa_paie_statut WHERE user_id=?', (uid,)).fetchone()[0] == 0


def test_b9_planning_absent_reste_en_attente(admin_client, db, sample_users):
    did = demande_partielle(db, sample_users['salarie_id'])
    approuver(admin_client, did)
    assert db.execute('SELECT statut FROM demandes_recup WHERE id=?', (did,)).fetchone()[0] == 'en_attente_direction'
    assert db.execute('SELECT COUNT(*) FROM heures_reelles').fetchone()[0] == 0


import pytest
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from tests.conftest import _login


TABLES = ('users', 'absences', 'demandes_conges', 'demandes_recup', 'heures_reelles',
          'presence_forfait_jour', 'validations', 'fiches_versions', 'fiches_evenements',
          'prepa_paie_statut', 'variables_paie', 'historique_modifications', 'journal_actions',
          'rh_projections', 'rh_projections_a_verifier', 'paie_a_recalculer')


def etat(db):
    return {t: [tuple(r) for r in db.execute(f'SELECT * FROM {t} ORDER BY rowid')] for t in TABLES}


@pytest.mark.parametrize('motif', ['Congé payé', 'Arrêt maladie', 'Congé conventionnel', 'Mi-temps thérapeutique'])
@pytest.mark.parametrize('profil', ['salarie_id', 'directeur_id'])
def test_b7_conflit_tous_motifs_atomique(admin_client, db, sample_users, motif, profil):
    uid = sample_users[profil]
    ajouter_absence(admin_client, uid)
    avant = etat(db)
    response = ajouter_absence(admin_client, uid, motif=motif)
    assert 'Conflit avec' in response.get_data(as_text=True)
    assert etat(db) == avant


def test_b7_maladie_maladie(admin_client, db, sample_users):
    uid = sample_users['salarie_id']
    ajouter_absence(admin_client, uid, motif='Arrêt maladie')
    avant = etat(db)
    ajouter_absence(admin_client, uid, motif='Arrêt maladie')
    assert etat(db) == avant


def test_b7_chevauchement_partiel_jours_1_a_8(admin_client, db, sample_users):
    uid = sample_users['salarie_id']
    ajouter_absence(admin_client, uid, '2026-09-01', '2026-09-05')
    avant = etat(db)
    ajouter_absence(admin_client, uid, '2026-09-04', '2026-09-08')
    assert etat(db) == avant
    for jour in range(1, 9):
        row = db.execute('SELECT type_saisie FROM heures_reelles WHERE user_id=? AND date=?', (uid, f'2026-09-{jour:02d}')).fetchone()
        assert bool(row) == (jour in (1, 2, 3, 4))
    assert db.execute('SELECT cp_pris FROM users WHERE id=?', (uid,)).fetchone()[0] == 4


def test_b7_chevauchement_deux_mois(admin_client, db, sample_users):
    uid = sample_users['salarie_id']
    ajouter_absence(admin_client, uid, '2026-08-31', '2026-09-02')
    avant = etat(db)
    ajouter_absence(admin_client, uid, '2026-09-02', '2026-09-04')
    assert etat(db) == avant
    assert db.execute('SELECT cp_pris FROM users WHERE id=?', (uid,)).fetchone()[0] == 3
    assert [r[0] for r in db.execute('SELECT date FROM heures_reelles ORDER BY date')] == ['2026-08-31','2026-09-01','2026-09-02']


def historique_absence(db, uid, auteur, motif, debut, fin, jours):
    row = db.execute('''INSERT INTO absences (user_id,motif,date_debut,date_fin,jours_ouvres,saisi_par)
        VALUES (?,?,?,?,?,?)''', (uid,motif,debut,fin,jours,auteur))
    return row.lastrowid


@pytest.mark.parametrize('motif', ['Congé payé', 'Arrêt maladie'])
def test_b7_suppression_historique_restaure_source_unique(admin_client, db, sample_users, motif):
    # Reproduction exacte d'une base antérieure, sans fabriquer de provenance.
    uid = sample_users['salarie_id']
    a = historique_absence(db,uid,sample_users['directeur_id'],'Congé payé','2026-09-01','2026-09-05',4)
    b = historique_absence(db,uid,sample_users['directeur_id'],motif,'2026-09-04','2026-09-08',3)
    from blueprints.absences import _reporter_absence_sur_calendrier
    _reporter_absence_sur_calendrier(db,a,uid,'2026-09-01','2026-09-05','Congé payé')
    _reporter_absence_sur_calendrier(db,b,uid,'2026-09-04','2026-09-08',motif)
    db.execute('DELETE FROM rh_projections')
    db.execute('UPDATE users SET cp_pris=? WHERE id=?', (7 if motif=='Congé payé' else 4,uid))
    db.commit()
    html = admin_client.get('/absences').get_data(as_text=True)
    assert 'Chevauchements historiques à examiner' in html
    admin_client.post(f'/absences/supprimer/{b}')
    assert db.execute('SELECT cp_pris FROM users WHERE id=?',(uid,)).fetchone()[0] == 4
    for jour in range(1,9):
        row = db.execute('SELECT commentaire FROM heures_reelles WHERE user_id=? AND date=?',(uid,f'2026-09-{jour:02d}')).fetchone()
        if jour in (1,2,3,4):
            assert row[0] == f'Absence #{a} - Congé payé'
        else:
            assert row is None
    assert db.execute('SELECT absence_id FROM rh_projections WHERE date="2026-09-04"').fetchone()[0] == a


def test_b7_suppression_ne_confond_pas_1_et_10(admin_client, db, sample_users):
    uid = sample_users['salarie_id']
    a = historique_absence(db,uid,sample_users['directeur_id'],'Congé payé','2026-09-07','2026-09-07',1)
    db.execute('''INSERT INTO absences (id,user_id,motif,date_debut,date_fin,jours_ouvres,saisi_par)
        VALUES (10,?,'Arrêt maladie','2026-09-07','2026-09-07',1,?)''', (uid,sample_users['directeur_id']))
    db.execute("INSERT INTO heures_reelles(user_id,date,type_saisie,commentaire,declaration_conforme) VALUES (?,'2026-09-07','absence','Absence #10 - Arrêt maladie',1)",(uid,))
    db.execute('UPDATE users SET cp_pris=1 WHERE id=?',(uid,))
    db.commit()
    admin_client.post(f'/absences/supprimer/{a}')
    assert db.execute('SELECT commentaire FROM heures_reelles').fetchone()[0] == 'Absence #10 - Arrêt maladie'


@pytest.mark.parametrize('ordre',['absence_puis_recup','recup_puis_absence'])
def test_b7_absence_et_recup_partielle_exclusives(admin_client, db, sample_users, sample_planning, ordre):
    uid = sample_users['salarie_id']
    did = demande_partielle(db,uid)
    if ordre == 'absence_puis_recup':
        ajouter_absence(admin_client,uid)
        avant = etat(db)
        approuver(admin_client,did)
    else:
        approuver(admin_client,did)
        avant = etat(db)
        ajouter_absence(admin_client,uid)
    assert etat(db) == avant


def test_b7_conge_approuve_origine_et_conflit(admin_client, db, sample_users):
    uid = sample_users['salarie_id']
    did = db.execute("INSERT INTO demandes_conges(user_id,type_conge,date_debut,date_fin,nb_jours) VALUES (?,'Congé payé','2026-09-07','2026-09-07',1)",(uid,)).lastrowid
    db.commit()
    approuver(admin_client,did,'conge')
    row = db.execute('SELECT * FROM absences').fetchone()
    assert row['demande_conge_id'] == did
    assert db.execute('SELECT absence_id FROM rh_projections').fetchone()[0] == row['id']
    avant = etat(db)
    ajouter_absence(admin_client,uid,motif='Arrêt maladie')
    assert etat(db) == avant


def test_b7_calendrier_ne_detruit_pas_projection(admin_client, db, sample_users, sample_contrat):
    uid = sample_users['salarie_id']
    ajouter_absence(admin_client,uid)
    avant = etat(db)
    response = admin_client.post('/saisie_heures',data={'user_id':uid,'date':'2026-09-07','type_saisie':'normal','heure_debut_matin':'08:30','heure_fin_matin':'12:00'},follow_redirects=True)
    assert response.status_code == 200
    assert etat(db) == avant


def reference_page(client, uid):
    html = client.get('/prepa_paie?mois=9&annee=2026').get_data(as_text=True)
    return re.search(r'name="reference_' + str(uid) + r'" value="([^"]+)"',html).group(1)


def test_b8_page_ancienne_refusee_autre_session(app, admin_client, db, sample_users, sample_contrat):
    uid=sample_users['salarie_id']
    ref=reference_page(admin_client,uid)
    compta=app.test_client()
    _login(compta,'compta_test','compta123')
    ajouter_absence(compta,uid,motif='Arrêt maladie')
    avant=etat(db)
    response=traiter(admin_client,uid,ref)
    assert 'Rechargez et vérifiez' in response.get_data(as_text=True)
    assert etat(db)==avant


@pytest.mark.parametrize('source',['contrat','piece_contrat','absence','piece_absence','variable','identite','secteur'])
def test_b8_sources_reellement_affichees(admin_client, db, sample_users, sample_contrat, source):
    uid=sample_users['salarie_id']
    ajouter_absence(admin_client,uid,motif='Arrêt maladie')
    traiter(admin_client,uid)
    ancien=dict(db.execute('SELECT * FROM prepa_paie_statut WHERE user_id=?',(uid,)).fetchone())
    mutations={
        'contrat':("UPDATE contrats SET date_fin='2026-09-20' WHERE user_id=?",uid),
        'piece_contrat':("UPDATE contrats SET fichier_path='piece-fictive.pdf' WHERE user_id=?",uid),
        'absence':("UPDATE absences SET date_reprise='2026-09-09' WHERE user_id=?",uid),
        'piece_absence':("UPDATE absences SET justificatif_path='piece-fictive.pdf' WHERE user_id=?",uid),
        'variable':("INSERT INTO variables_paie(user_id,mois,annee,transport) VALUES (?,9,2026,17)",uid),
        'identite':("UPDATE users SET prenom='Jean-Paul' WHERE id=?",uid),
        'secteur':("UPDATE secteurs SET nom='Secteur Renommé' WHERE id=?",sample_users['secteur_id']),
    }
    sql,param=mutations[source];db.execute(sql,(param,));db.commit()
    statut=db.execute('SELECT * FROM prepa_paie_statut WHERE user_id=?',(uid,)).fetchone()
    assert statut['traite']==0 and statut['modifie_le']
    assert statut['verifie_le']==ancien['verifie_le']
    assert statut['empreinte_verifiee']==ancien['empreinte_verifiee']


@pytest.mark.parametrize('source',['email','planning','heures','autre_mois','autre_salarie','document_hors_grille','contrat_hors_mois'])
def test_b8_modification_hors_donnees_preserve_verification(admin_client, db, sample_users, sample_contrat, source):
    uid=sample_users['salarie_id']
    traiter(admin_client,uid)
    avant=dict(db.execute('SELECT * FROM prepa_paie_statut WHERE user_id=?',(uid,)).fetchone())
    mutations={
        'email':("UPDATE users SET email='fictif@example.test' WHERE id=?",uid),
        'planning':("INSERT INTO planning_theorique(user_id,type_periode,date_debut_validite) VALUES (?,'periode_scolaire','2026-01-01')",uid),
        'heures':("INSERT INTO heures_reelles(user_id,date,type_saisie) VALUES (?,'2026-09-08','normal')",uid),
        'autre_mois':("INSERT INTO variables_paie(user_id,mois,annee,transport) VALUES (?,10,2026,17)",uid),
        'autre_salarie':("INSERT INTO variables_paie(user_id,mois,annee,transport) VALUES (?,9,2026,17)",sample_users['responsable_id']),
        'document_hors_grille':("INSERT INTO documents_salaries(user_id,type_document,fichier_path,fichier_nom) VALUES (?,'autre','fictif.pdf','fictif.pdf')",uid),
        'contrat_hors_mois':("INSERT INTO contrats(user_id,type_contrat,date_debut,date_fin) VALUES (?,'CDD','2025-01-01','2025-02-01')",uid),
    }
    sql,param=mutations[source];db.execute(sql,(param,));db.commit()
    assert dict(db.execute('SELECT * FROM prepa_paie_statut WHERE user_id=?',(uid,)).fetchone())==avant


def test_b8_reference_autre_dossier_et_absente_refusees(admin_client, db, sample_users, sample_contrat):
    uid=sample_users['salarie_id'];ref=reference_page(admin_client,uid)
    for cible,jeton in [(uid,''),(sample_users['responsable_id'],ref),(uid,ref+'x')]:
        avant=etat(db);traiter(admin_client,cible,jeton);assert etat(db)==avant


def test_b9_reprise_planning_et_notification_une_fois(admin_client, db, sample_users, monkeypatch):
    uid=sample_users['salarie_id'];did=demande_partielle(db,uid)
    appels=[]
    monkeypatch.setattr('blueprints.recup.is_email_configured',lambda: True)
    monkeypatch.setattr('blueprints.recup.peut_envoyer_email',lambda uid:(True,'fictif@example.test'))
    monkeypatch.setattr('blueprints.recup.notifier_demande_recup_decision',lambda *a,**kw:appels.append(a))
    avant=etat(db)
    response=approuver(admin_client,did)
    assert 'planning absent' in response.get_data(as_text=True)
    assert etat(db)==avant and not appels
    db.execute("""INSERT INTO planning_theorique(user_id,type_periode,date_debut_validite,lundi_matin_debut,lundi_matin_fin,lundi_aprem_debut,lundi_aprem_fin)
        VALUES (?,'periode_scolaire','2000-01-01','08:30','12:00','13:30','17:00')""",(uid,));db.commit()
    approuver(admin_client,did)
    apres=etat(db)
    approuver(admin_client,did)
    assert etat(db)==apres and len(appels)==1
    assert db.execute('SELECT demande_recup_id FROM rh_projections').fetchone()[0]==did
    assert db.execute("SELECT COUNT(*) FROM historique_modifications WHERE action='application_recup'").fetchone()[0]==1


def test_b9_approbations_concurrentes_une_application(app, db, sample_users, sample_planning):
    did=demande_partielle(db,sample_users['salarie_id'])
    clients=[app.test_client(),app.test_client()]
    for c in clients:_login(c,'admin','Admin1234')
    depart=Barrier(2)
    def action(c):
        depart.wait(timeout=10)
        return approuver(c,did).status_code
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert list(pool.map(action,clients))==[200,200]
    assert db.execute('SELECT COUNT(*) FROM rh_projections').fetchone()[0]==1
    assert db.execute("SELECT COUNT(*) FROM historique_modifications WHERE action='application_recup'").fetchone()[0]==1


@pytest.mark.parametrize('route',['absences','prepa','recup'])
def test_droits_revoques_aucune_mutation(app, db, sample_users, sample_contrat, route):
    c=app.test_client();_login(c,'admin','Admin1234')
    uid=sample_users['salarie_id'];ref=reference_page(c,uid);did=demande_partielle(db,uid)
    db.execute('UPDATE users SET session_version=session_version+1 WHERE id=?',(sample_users['directeur_id'],));db.commit()
    avant=etat(db)
    if route=='absences':ajouter_absence(c,uid)
    elif route=='prepa':traiter(c,uid,ref)
    else:approuver(c,did)
    assert etat(db)==avant


def test_b9_journee_sans_planning_atomique(admin_client,db,sample_users):
    uid=sample_users['salarie_id']
    did=db.execute("INSERT INTO demandes_recup(user_id,date_debut,date_fin,nb_jours,nb_heures,statut) VALUES (?,'2026-09-07','2026-09-07',1,7,'en_attente_direction')",(uid,)).lastrowid
    db.commit();avant=etat(db)
    response=approuver(admin_client,did)
    assert 'planning absent' in response.get_data(as_text=True)
    assert etat(db)==avant


@pytest.mark.parametrize('debut,fin',[('invalide','2026-09-07'),('2026-09-07','2026-02-30'),('2026-09-08','2026-09-07')])
def test_b7_dates_invalides_refusees(admin_client,db,sample_users,debut,fin):
    avant=etat(db)
    assert ajouter_absence(admin_client,sample_users['salarie_id'],debut,fin).status_code==200
    assert etat(db)==avant


def test_b7_dates_normalisees_avant_conflit(admin_client,db,sample_users):
    uid=sample_users['salarie_id'];ajouter_absence(admin_client,uid)
    avant=etat(db);ajouter_absence(admin_client,uid,'2026-9-7');assert etat(db)==avant


def test_b8_page_ancienne_ne_decoche_pas_verification_concurrente(admin_client,db,sample_users,sample_contrat):
    uid=sample_users['salarie_id'];ref=reference_page(admin_client,uid)
    traiter(admin_client,uid)
    avant=etat(db)
    admin_client.post('/prepa_paie/traiter',data={'mois':9,'annee':2026,'user_ids':[uid],f'reference_{uid}':ref})
    assert etat(db)==avant


def test_b8_export_affiche_obsolescence(admin_client,db,sample_users,sample_contrat):
    from io import BytesIO
    from openpyxl import load_workbook
    uid=sample_users['salarie_id'];traiter(admin_client,uid);ajouter_absence(admin_client,uid)
    response=admin_client.get('/prepa_paie/export_excel?mois=9&annee=2026')
    ws=load_workbook(BytesIO(response.data)).active
    assert ws.cell(row=2,column=1).value=='Modifié depuis la dernière vérification'


def test_b7_conge_refuse_sur_absence_deja_presente(admin_client,db,sample_users,monkeypatch):
    uid=sample_users['salarie_id'];ajouter_absence(admin_client,uid)
    did=db.execute("INSERT INTO demandes_conges(user_id,type_conge,date_debut,date_fin,nb_jours) VALUES (?,'Congé payé','2026-09-07','2026-09-07',1)",(uid,)).lastrowid
    db.commit();avant=etat(db);appels=[]
    monkeypatch.setattr('blueprints.recup.is_email_configured',lambda:appels.append(True) or False)
    approuver(admin_client,did,'conge')
    assert etat(db)==avant and not appels


def test_b7_calendrier_forfait_protege_origine(admin_client,db,sample_users):
    uid=sample_users['directeur_id'];ajouter_absence(admin_client,uid,motif='Arrêt maladie')
    avant=etat(db)
    response=admin_client.post('/calendrier_forfait_jour',data={'date':'2026-09-07','type_journee':'travaille'},follow_redirects=True)
    assert response.status_code==200
    assert etat(db)==avant


def test_b7_projection_historique_non_ecrasable(admin_client,db,sample_users,sample_contrat):
    uid=sample_users['salarie_id'];ajouter_absence(admin_client,uid)
    db.execute('DELETE FROM rh_projections');db.commit()
    avant=etat(db)
    admin_client.post('/saisie_heures',data={'user_id':uid,'date':'2026-09-07','type_saisie':'normal','heure_debut_matin':'08:30','heure_fin_matin':'12:00'})
    assert etat(db)==avant


def test_b7_auto_validation_conge_direction_refuse_conflit(admin_client,db,sample_users):
    uid=sample_users['directeur_id'];ajouter_absence(admin_client,uid,motif='Arrêt maladie')
    avant=etat(db)
    response=admin_client.post('/demande_conge',data={'type_conge':'Congé payé','date_debut':'2026-09-07','date_fin':'2026-09-07'},follow_redirects=True)
    assert 'Conflit avec' in response.get_data(as_text=True)
    assert etat(db)==avant


def test_b8_formulaire_plusieurs_salaries_rollback_global(admin_client,db,sample_users,sample_contrat):
    uid=sample_users['salarie_id'];autre=sample_users['comptable_id']
    db.execute("INSERT INTO contrats(user_id,type_contrat,date_debut) VALUES (?,'CDI','2026-01-01')",(autre,));db.commit()
    refs={cible:reference_page(admin_client,cible) for cible in (autre,uid)}
    ajouter_absence(admin_client,uid,motif='Arrêt maladie')
    avant=etat(db)
    admin_client.post('/prepa_paie/traiter',data={'mois':9,'annee':2026,'user_ids':[autre,uid],
        f'traite_{autre}':'1',f'traite_{uid}':'1',f'reference_{autre}':refs[autre],f'reference_{uid}':refs[uid]})
    assert etat(db)==avant


def test_b7_absence_et_recuperation_concurrentes(app,db,sample_users,sample_planning):
    uid=sample_users['salarie_id'];did=demande_partielle(db,uid)
    clients=[app.test_client(),app.test_client()]
    for c in clients:_login(c,'admin','Admin1234')
    depart=Barrier(2)
    def action(n):
        depart.wait(timeout=10)
        return (ajouter_absence(clients[n],uid) if n==0 else approuver(clients[n],did)).status_code
    with ThreadPoolExecutor(max_workers=2) as pool:assert list(pool.map(action,(0,1)))==[200,200]
    nb_absences=db.execute('SELECT COUNT(*) FROM absences').fetchone()[0]
    valide=db.execute("SELECT COUNT(*) FROM demandes_recup WHERE statut='validee'").fetchone()[0]
    assert nb_absences+valide==1
    assert db.execute('SELECT COUNT(*) FROM rh_projections').fetchone()[0]==1
    assert db.execute('SELECT cp_pris FROM users WHERE id=?',(uid,)).fetchone()[0]==nb_absences


def test_b7_post_conge_date_non_normalisee_ne_contourne_pas_conflit(admin_client,db,sample_users):
    uid=sample_users['directeur_id'];ajouter_absence(admin_client,uid,motif='Arrêt maladie')
    avant=etat(db)
    response=admin_client.post('/demande_conge',data={'type_conge':'Congé payé','date_debut':'2026-9-7','date_fin':'2026-9-7'},follow_redirects=True)
    assert 'Conflit avec' in response.get_data(as_text=True)
    assert etat(db)==avant


def test_b7_demande_historique_date_invalide_reste_en_attente(admin_client,db,sample_users):
    did=db.execute("INSERT INTO demandes_conges(user_id,type_conge,date_debut,date_fin,nb_jours) VALUES (?,'Congé payé','2026-9-7','2026-9-7',1)",(sample_users['salarie_id'],)).lastrowid
    db.commit();avant=etat(db)
    response=approuver(admin_client,did,'conge')
    assert 'Dates de la demande' in response.get_data(as_text=True)
    assert etat(db)==avant
